"""Excel spread with live ratio formulas, plus headless LibreOffice recalculation.

Sheets:
* Spread    - standardized line items by fiscal year, in $ millions (inputs)
* Ratios    - every credit ratio as a formula over Spread cells
* Covenants - latest-year covenant tests: value (linked to Ratios), threshold,
              headroom and PASS/BREACH as formulas
* Stress    - breaking shock per scenario and covenant (values)
* Sources   - the XBRL tag behind each spread cell

Leverage ratios show "n/m" (not meaningful) when EBITDA <= 0, matching the
infinite values on the Python side.
"""

from __future__ import annotations

import math
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from credit_engine.ratios import AMOUNT_RATIOS, RATIO_LABELS
from credit_engine.spreading import LINE_ITEMS, Spread

SCALE = 1e6
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FONT = Font(color="0000FF")  # banking convention: blue = hard-coded input

# ratio -> formula template using {item} placeholders resolved to Spread cells
RATIO_FORMULAS = {
    "ebitda": "{operating_income}+{depreciation_amortization}+{impairments}",
    "total_debt": "{short_term_borrowings}+{current_portion_ltd}+{long_term_debt}",
    "net_debt": "{short_term_borrowings}+{current_portion_ltd}+{long_term_debt}-{cash}",
    "ebitda_margin": "IF({revenue}=0,\"n/m\",{EBITDA}/{revenue})",
    "debt_to_ebitda": "IF({EBITDA}<=0,\"n/m\",{TOTAL_DEBT}/{EBITDA})",
    "net_leverage": "IF({EBITDA}<=0,\"n/m\",{NET_DEBT}/{EBITDA})",
    "interest_coverage": "IF({interest_expense}=0,\"n/m\",({operating_income}+{impairments})/{interest_expense})",
    "dscr": "IF(({interest_expense}+{current_portion_ltd})=0,\"n/m\",{EBITDA}/({interest_expense}+{current_portion_ltd}))",
    "fccr": "IF(({interest_expense}+{operating_lease_cost}+{current_portion_ltd})=0,\"n/m\","
            "({EBITDA}+{operating_lease_cost}-{capex}-{cash_taxes})"
            "/({interest_expense}+{operating_lease_cost}+{current_portion_ltd}))",
    "current_ratio": "IF({current_liabilities}=0,\"n/m\",{current_assets}/{current_liabilities})",
    "free_cash_flow": "{cfo}-{capex}",
    "debt_to_equity": "IF({total_equity}<=0,\"n/m\",{TOTAL_DEBT}/{total_equity})",
}


def _header(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL


def build_workbook(spread: Spread, ratios_by_year: dict[int, dict], covenants: list, stress: list,
                   path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Spread"
    years = spread.years
    _header(ws, ["Line item ($M)", "Statement"] + [f"FY{y}" for y in years])
    item_row = {}
    for item, (stmt, *_rest) in LINE_ITEMS.items():
        vals = [spread.values.at[item, y] for y in years]
        ws.append([item, stmt] + [None if math.isnan(v) else v / SCALE for v in vals])
        item_row[item] = ws.max_row
        for c in ws[ws.max_row][2:]:
            c.font = INPUT_FONT
            c.number_format = "#,##0.0"
    ws.append([])
    ws.append(["Period end", ""] + [spread.period_ends[y] for y in years])
    ws.column_dimensions["A"].width = 28

    wr = wb.create_sheet("Ratios")
    _header(wr, ["Ratio", "Key"] + [f"FY{y}" for y in years])
    ratio_row = {}
    for i, key in enumerate(RATIO_FORMULAS):
        ratio_row[key] = i + 2
    for key, tmpl in RATIO_FORMULAS.items():
        row = [RATIO_LABELS[key], key]
        for j, _y in enumerate(years):
            col = get_column_letter(3 + j)
            refs = {item: f"Spread!{col}{r}" for item, r in item_row.items()}
            refs.update({"EBITDA": f"{col}{ratio_row['ebitda']}", "TOTAL_DEBT": f"{col}{ratio_row['total_debt']}",
                         "NET_DEBT": f"{col}{ratio_row['net_debt']}"})
            row.append("=" + tmpl.format(**refs))
        wr.append(row)
        for c in wr[wr.max_row][2:]:
            c.number_format = "#,##0.0" if key in AMOUNT_RATIOS else ("0.0%" if key == "ebitda_margin" else "0.00x")
    wr.column_dimensions["A"].width = 28

    wc = wb.create_sheet("Covenants")
    last_col = get_column_letter(2 + len(years))
    _header(wc, ["Covenant", "Metric", "Test", "Threshold", f"Actual FY{years[-1]}", "Headroom", "Status"])
    for c in covenants:
        r = wc.max_row + 1
        actual = f"Ratios!{last_col}{ratio_row[c.metric]}"
        if c.op == "<=":
            head = f'=IF(ISNUMBER(E{r}),(D{r}-E{r})/D{r},"n/m")'
        else:
            head = f'=IF(ISNUMBER(E{r}),(E{r}-D{r})/D{r},"n/m")'
        wc.append([c.name, c.metric, c.op, c.threshold, f"={actual}", head,
                   f'=IF(ISNUMBER(F{r}),IF(F{r}>=0,"PASS","BREACH"),IF(C{r}="<=","BREACH","PASS"))'])
        wc[f"D{r}"].font = INPUT_FONT
        wc[f"F{r}"].number_format = "0.0%"
    wc.column_dimensions["A"].width = 28

    wsx = wb.create_sheet("Stress")
    _header(wsx, ["Scenario", "Covenant", "Breaking shock", "Result"])
    for s in stress:
        wsx.append([s.scenario, s.covenant, s.breaking_shock, s.status])

    wsrc = wb.create_sheet("Sources")
    _header(wsrc, ["Line item"] + [f"FY{y}" for y in years])
    for item in LINE_ITEMS:
        wsrc.append([item] + [spread.sources[(item, y)] for y in years])
    wsrc.column_dimensions["A"].width = 28

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def soffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def recalculate(path: Path, timeout: int = 180) -> Path:
    """Open the workbook in headless LibreOffice, recalculate and save a copy; returns the copy's path."""
    exe = soffice_path()
    if exe is None:
        raise RuntimeError("LibreOffice (soffice) not found; install libreoffice-calc")
    path = Path(path).resolve()
    with tempfile.TemporaryDirectory(prefix="credit-engine-lo-") as profile:
        outdir = path.parent / "recalc"
        outdir.mkdir(exist_ok=True)
        cmd = [exe, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
               "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(outdir), str(path)]
        subprocess.run(cmd, check=True, capture_output=True, timeout=timeout)
    out = outdir / path.name
    if not out.exists():
        raise RuntimeError(f"LibreOffice did not write {out}")
    return out


def read_ratio_values(path: Path) -> dict[int, dict[str, object]]:
    """Read computed ratio values (cached by LibreOffice) keyed by fiscal year."""
    wb = load_workbook(path, data_only=True)
    ws = wb["Ratios"]
    header = [c.value for c in ws[1]]
    out: dict[int, dict[str, object]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = row[1]
        for h, v in zip(header[2:], row[2:]):
            out.setdefault(int(str(h)[2:]), {})[key] = v
    return out


def read_covenant_values(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb["Covenants"]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def compare_with_python(recalc_path: Path, ratios_by_year: dict[int, dict], rel_tol: float = 1e-9) -> list[str]:
    """Return a list of mismatches between workbook formulas and Python ratios (empty = parity)."""
    excel = read_ratio_values(recalc_path)
    problems = []
    for fy, ratios in ratios_by_year.items():
        for key, py in ratios.items():
            xv = excel[fy][key]
            if key in AMOUNT_RATIOS:
                py = py / SCALE
            if math.isinf(py) or math.isnan(py):
                if xv != "n/m":
                    problems.append(f"FY{fy} {key}: python {py}, excel {xv!r}")
            elif not isinstance(xv, (int, float)) or not math.isclose(xv, py, rel_tol=rel_tol, abs_tol=1e-9):
                problems.append(f"FY{fy} {key}: python {py}, excel {xv!r}")
    return problems
