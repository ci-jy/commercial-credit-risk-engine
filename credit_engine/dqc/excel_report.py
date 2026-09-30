"""Excel exceptions workbook for a screened quarter.

Sheets:
* Summary   - findings and filings flagged per rule; each rule ID links to its sheet
* DQC_XXXX  - one sheet per rule: filer, CIK, form, accession (linked to EDGAR),
              concept, period, dimensions, reported value, suggested correction and
              the rule message; cell A1 links back to the summary
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from credit_engine.dqc import RULE_IDS, rule_module

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
LINK_FONT = Font(color="0563C1", underline="single")
EDGAR_INDEX = "https://www.sec.gov/Archives/edgar/data/{cik}/{nodash}/{adsh}-index.htm"
COLUMNS = ["Filer", "CIK", "Form", "Accession", "Concept", "Period end", "Quarters", "Dimensions", "Unit",
           "Reported value", "Suggested correction", "Rule message"]
WIDTHS = [34, 10, 7, 22, 46, 11, 9, 40, 8, 18, 20, 90]


def _header(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL


def build_exceptions_workbook(findings: pd.DataFrame, sub: pd.DataFrame, path: Path, title: str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    meta = sub.set_index("adsh")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{len(findings):,} findings in {findings.adsh.nunique():,} of {len(sub):,} filings"])
    ws.append([])
    _header(ws, ["Rule", "Title", "Findings", "Filings flagged", "Values with a suggested correction"])
    for rid in RULE_IDS:
        f = findings[findings.rule == rid]
        ws.append([rid, rule_module(rid).TITLE, len(f), f.adsh.nunique(), int(f.suggested.notna().sum())])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.hyperlink = f"#'{rid}'!A1"
        cell.font = LINK_FONT
    for i, w in enumerate([12, 60, 10, 15, 32], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for rid in RULE_IDS:
        rs = wb.create_sheet(rid)
        rs.append(["<< Summary", f"{rid}: {rule_module(rid).TITLE}"])
        rs["A1"].hyperlink = "#'Summary'!A1"
        rs["A1"].font = LINK_FONT
        rs["B1"].font = Font(bold=True)
        _header(rs, COLUMNS)
        f = findings[findings.rule == rid]
        for r in f.itertuples():
            m = meta.loc[r.adsh] if r.adsh in meta.index else None
            cik = int(m.cik) if m is not None else 0
            rs.append([
                m["name"] if m is not None else "", cik, m.form if m is not None else "", r.adsh,
                r.concept, pd.Timestamp(str(r.ddate)).date(), int(r.qtrs), r.segments or "", r.uom,
                float(r.value), None if pd.isna(r.suggested) else float(r.suggested), r.message,
            ])
            acc = rs.cell(row=rs.max_row, column=4)
            acc.hyperlink = EDGAR_INDEX.format(cik=cik, nodash=r.adsh.replace("-", ""), adsh=r.adsh)
            acc.font = LINK_FONT
            rs.cell(row=rs.max_row, column=6).number_format = "yyyy-mm-dd"
            rs.cell(row=rs.max_row, column=10).number_format = "#,##0.####"
            rs.cell(row=rs.max_row, column=11).number_format = "#,##0.####"
        for i, w in enumerate(WIDTHS, start=1):
            rs.column_dimensions[get_column_letter(i)].width = w
        rs.freeze_panes = "A3"
        if f.shape[0]:
            rs.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{rs.max_row}"
    wb.save(path)
    return path
