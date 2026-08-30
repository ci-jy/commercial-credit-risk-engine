"""Excel formula parity: workbook formulas recalculated by headless LibreOffice equal the Python ratios."""

import math

import pytest
from openpyxl import load_workbook

from credit_engine.edgar import load_borrowers, load_fixture
from credit_engine.excel import (build_workbook, compare_with_python, read_covenant_values, read_ratio_values,
                                 recalculate, soffice_path)
from credit_engine.ratios import DEFAULT_COVENANTS, compute_ratios, test_covenants as run_covenants
from credit_engine.spreading import spread_companyfacts
from credit_engine.stress import run_stress

needs_lo = pytest.mark.skipif(soffice_path() is None, reason="LibreOffice (soffice) is not installed")


def _build(name, tmp_path, covenants=None):
    s = spread_companyfacts(load_fixture(name))
    covenants = covenants or DEFAULT_COVENANTS
    rb = {fy: compute_ratios(s.year(fy)) for fy in s.years}
    cov = run_covenants(rb[s.latest_year], covenants)
    path = build_workbook(s, rb, cov, run_stress(s.year(), covenants, 0.5), tmp_path / f"{name}.xlsx")
    return s, rb, cov, path


def test_ratio_cells_are_formulas(tmp_path):
    _, _, _, path = _build("sample_borrower", tmp_path)
    ws = load_workbook(path)["Ratios"]
    cells = [c.value for row in ws.iter_rows(min_row=2, min_col=3) for c in row]
    assert cells and all(isinstance(v, str) and v.startswith("=") for v in cells)
    assert "Spread!" in ws["C2"].value


@needs_lo
@pytest.mark.parametrize("name", ["sample_borrower", "hasbro", "whirlpool"])
def test_libreoffice_recalc_matches_python(name, tmp_path):
    cov_spec = load_borrowers()[name].get("covenants", DEFAULT_COVENANTS)
    s, rb, cov, path = _build(name, tmp_path, cov_spec)
    recalced = recalculate(path)
    assert compare_with_python(recalced, rb) == []
    values = read_ratio_values(recalced)
    assert len(values) == len(s.years)
    if name == "sample_borrower":
        assert values[2024]["debt_to_ebitda"] == pytest.approx(3.0, rel=1e-12)
        assert values[2024]["fccr"] == pytest.approx(40 / 29, rel=1e-12)
    for c, row in zip(cov, read_covenant_values(recalced)):
        assert row["Status"] == ("PASS" if c.passed else "BREACH")
        if not math.isinf(c.headroom):
            assert row["Headroom"] == pytest.approx(c.headroom, rel=1e-9, abs=1e-12)


@needs_lo
def test_negative_ebitda_shows_nm(tmp_path):
    # force negative EBITDA in one year: leverage cells must read "n/m" like the infinite Python values
    s = spread_companyfacts(load_fixture("sample_borrower"))
    s.values.loc["operating_income", 2023] = -30e6
    rb = {fy: compute_ratios(s.year(fy)) for fy in s.years}
    path = build_workbook(s, rb, run_covenants(rb[2024]), [], tmp_path / "neg.xlsx")
    recalced = recalculate(path)
    assert compare_with_python(recalced, rb) == []
    assert read_ratio_values(recalced)[2023]["debt_to_ebitda"] == "n/m"
