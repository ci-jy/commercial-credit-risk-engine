"""Ratio impact, the Excel exceptions workbook and the bulk-vs-Arelle comparison."""

import math

import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import ADSH, build_quarter, fact
from credit_engine import REPO_ROOT
from credit_engine.dqc import RULE_IDS, run_screen
from credit_engine.dqc.compare import (CATEGORIES, DQC_FIXTURES, agreement_by_rule, align, explain, load_arelle,
                                       load_root_causes)
from credit_engine.dqc.excel_report import build_exceptions_workbook
from credit_engine.dqc.fsds import load_quarter
from credit_engine.dqc.impact import ratio_impact, summarise_impact


def spread_facts(debt=-3_000_000):
    return [fact("OperatingIncomeLoss", 1_000_000, qtrs=4), fact("DepreciationDepletionAndAmortization", 200_000, qtrs=4),
            fact("InterestExpense", 100_000, qtrs=4), fact("LongTermDebtNoncurrent", debt),
            fact("AssetsCurrent", 2_000_000), fact("LiabilitiesCurrent", 1_000_000),
            fact("StockholdersEquity", 4_000_000)]


def test_sign_flipped_debt_changes_leverage():
    q = build_quarter(spread_facts())
    findings = run_screen(q)
    debt = findings[findings.concept == "LongTermDebtNoncurrent"]
    assert debt.rule.tolist() == ["DQC_0015"] and debt.suggested.iat[0] == 3_000_000
    impact = ratio_impact(findings, q.num, q.sub)
    lev = impact[(impact.concept == "LongTermDebtNoncurrent") & (impact.ratio == "debt_to_ebitda")].iloc[0]
    assert lev.before == pytest.approx(-2.5) and lev.after == pytest.approx(2.5) and lev.changed
    cur = impact[impact.ratio == "current_ratio"].iloc[0]
    assert not cur.changed
    summ = summarise_impact(impact).set_index("ratio")
    assert summ.loc["debt_to_ebitda", "changed"] == 1 and summ.loc["debt_to_ebitda", "median_abs_pct"] == pytest.approx(200)


def test_clean_spread_has_no_impact():
    q = build_quarter(spread_facts(debt=3_000_000))
    assert ratio_impact(run_screen(q), q.num, q.sub).empty


def test_quarterly_filing_annualises_flows():
    facts = [dict(f, qtrs=1) if f["qtrs"] == 4 else f for f in spread_facts()]
    q = build_quarter(facts, form="10-Q")
    impact = ratio_impact(run_screen(q), q.num, q.sub)
    lev = impact[impact.ratio == "debt_to_ebitda"].iloc[0]
    assert lev.after == pytest.approx(3_000_000 / (4 * 1_200_000))


def _check_links(path):
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", *RULE_IDS]
    internal = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                link = cell.hyperlink
                if link is None:
                    continue
                target = link.location or link.target
                if target.startswith("http"):
                    assert target.startswith("https://www.sec.gov/Archives/edgar/data/")
                    continue
                sheet, ref = target.lstrip("#").rsplit("!", 1)
                assert sheet.strip("'") in wb.sheetnames and ref == "A1"
                internal += 1
    return wb, internal


def test_exceptions_workbook_links_resolve(tmp_path):
    q = build_quarter(spread_facts() + [fact("Goodwill", -5_000_000)])
    findings = run_screen(q)
    path = build_exceptions_workbook(findings, q.sub, tmp_path / "x.xlsx", "test")
    wb, internal = _check_links(path)
    assert internal == 2 * len(RULE_IDS)  # summary -> rule sheet, rule sheet -> summary
    ws = wb["DQC_0015"]
    assert ws["A3"].value == "PLANTED CO" and ws["D3"].value == ADSH and ws["L3"].value == 3_000_000
    summary = {r[0]: r[2] for r in wb["Summary"].iter_rows(min_row=5, values_only=True)}
    assert summary["DQC_0015"] == 1 and summary["DQC_0014"] == 1


def test_committed_quarter_workbook_opens_and_links_resolve():
    path = REPO_ROOT / "reports" / "dqc_exceptions_2026q2.xlsx"
    if not path.exists():
        pytest.skip("full-quarter workbook not generated")
    _check_links(path)


# Differential comparison -------------------------------------------------------------------------

@pytest.fixture(scope="module")
def comparison():
    sample = pd.read_csv(DQC_FIXTURES / "sample_filings.csv", dtype=str)
    arelle, runs = load_arelle(list(sample.adsh))
    q = load_quarter(DQC_FIXTURES / "fsds_sample", adsh=set(runs.adsh))
    return explain(align(run_screen(q), arelle)), runs


def test_reference_runs_cover_40_plus_filings(comparison):
    _, runs = comparison
    assert len(runs) >= 40
    assert (runs.runtime_s > 0).all() and runs.ruleset.str.startswith("dqc-us-20").all()


def test_every_disagreement_is_root_caused(comparison):
    explained, _ = comparison
    one_sided = explained[explained.side != "both"]
    assert (one_sided.category != "unexplained").all(), one_sided[one_sided.category == "unexplained"].key.tolist()
    assert set(one_sided.category) <= set(CATEGORIES)
    assert (one_sided.root_cause.str.len() > 20).all()


def test_root_cause_entries_are_all_used(comparison):
    explained, _ = comparison
    one_sided = explained[explained.side != "both"].to_dict("records")
    from credit_engine.dqc.compare import _matches
    for rc in load_root_causes():
        assert any(_matches(r, rc["match"]) for r in one_sided), rc


def test_align_and_agreement_on_synthetic_findings():
    bulk = pd.DataFrame([
        {"adsh": "a", "rule": "DQC_0015", "element_id": "1", "concept": "Assets", "ddate": 20251231, "qtrs": 0,
         "segments": "", "uom": "USD", "value": -1.0, "suggested": 1.0, "message": ""},
        {"adsh": "a", "rule": "DQC_0015", "element_id": "1", "concept": "Liabilities", "ddate": 20251231, "qtrs": 0,
         "segments": "B=y;A=x;", "uom": "USD", "value": -1.0, "suggested": 1.0, "message": ""}])
    arelle = pd.DataFrame([
        {"adsh": "a", "rule": "DQC_0015", "element_id": "1", "concept": "Assets", "ddate": 20251231, "segments": "",
         "value": "-1", "message": ""},
        {"adsh": "a", "rule": "DQC_0015", "element_id": "1", "concept": "Liabilities", "ddate": 20251231,
         "segments": "A=x;B=y;", "value": "-1", "message": ""},
        {"adsh": "b", "rule": "DQC_0004", "element_id": "16", "concept": "Assets", "ddate": 20251231, "segments": "",
         "value": "5", "message": ""}])
    m = explain(align(bulk, arelle), root_causes=[])
    assert sorted(m.side) == ["arelle", "both", "both"]
    agree = agreement_by_rule(m, n_filings=2).set_index("rule")
    assert agree.loc["DQC_0015", "finding_agreement"] == 1.0
    assert agree.loc["DQC_0004", "finding_agreement"] == 0.0 and agree.loc["DQC_0004", "filing_agreement"] == 0.5
    assert agree.loc["DQC_0004", "unexplained"] == 1
    assert math.isclose(agree.loc["DQC_0001", "finding_agreement"], 1.0)
